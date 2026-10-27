#!/usr/bin/env python3
"""build_pptx.py -- host-neutral PowerPoint build/validate/render for the deck skill.

Uses python-pptx (structural build + reopen-and-check) and, when present, a LibreOffice
`soffice` binary (structural render to PNG per slide) for visual inspection. Neither is
guaranteed to be installed on every host -- run --health-check first and report the exact
missing piece rather than attempting an unvalidated build.

    python build_pptx.py --health-check
    python build_pptx.py --build SPEC.json OUT.pptx [--template TPL.pptx]
    python build_pptx.py --profile-template TPL.pptx
    python build_pptx.py --validate OUT.pptx [--template TPL.pptx]
    python build_pptx.py --render OUT.pptx OUT_DIR

SPEC.json shape:
    {"title": "Deck title",
     "title_layout": 0,                              # optional index or layout name
     "slides": [{"title": "Slide title", "bullets": ["...", "..."], "layout": 1}, ...]}

A bullet entry is normally a plain string. It may instead be `{"runs": [...]}` to mix body text
with an editor-facing note -- a source citation/locator or a `voice.md` -> Uncertainty marker
`[verify]` flag -- so the note reads as visibly distinct from body text:
    {"runs": [{"text": "Effect size was 4.2 points "},
               {"text": "[Smith 2024, p.12]", "note": "reference"}]}
Each run is `{"text": "..."}`, optionally with `"note"`: `"reference"` (teal), `"primary-reference"`
(purple, the source-native marker a `reference` claim traces back to) or `"certainty"` (amber) --
the same colors and the same font-color-plus-italic mechanism build_docx.py uses, so a note reads
the same way in a deck as it does in a Word document. A claim with both layers of provenance gets
both runs, back to back -- see deck-rules.md -> Editor notes and docx-rules.md -> Editor notes ->
Dual citation for the exact shape and when it applies. The marker text itself is unchanged by
this -- only its formatting is. An unrecognized note type is refused, not guessed at.

`layout` (per slide) and `title_layout` accept either a layout index (int) or a layout name
(str, e.g. "Title and Content") -- run --profile-template first to see what a given template
actually offers; deck-rules.md requires this before building against an unfamiliar template
rather than guessing indices from visual similarity. Indices cover every slide master's layouts
flattened together, not just the first master's -- a template with more than one master (a
common shape for a distributed corporate deck: one master for covers, another for content) has
layouts past the first master that would otherwise be invisible to this tool. Without --template,
a blank default presentation's own layouts 0/1 are used, matching prior behavior.

A layout --profile-template reports with `"usable": false` has two or more placeholders sharing
both idx and type -- python-pptx and PowerPoint can't tell them apart, so which one a slide's
content actually lands in isn't guaranteed. --build refuses to use such a layout rather than
risk silently ambiguous output; pick a different layout, or fix the template in PowerPoint's
Slide Master view (give each placeholder a distinct idx).

--template accepts a .pptx, .potx, or .ppsx file. It uses that presentation's master, theme,
dimensions, and slide layouts instead of a blank default deck, per deck-rules.md's template
precedence. Existing template slides are removed: the default is a new presentation in that
visual system, not a copy of another project's content. The template file itself is never
modified; the output must be a new file. `--validate` with `--template` additionally checks the
built deck kept the template's slide dimensions.

Structural --validate is always available wherever --build is (same python-pptx dependency).
--render additionally needs a `soffice` (LibreOffice) binary (on PATH or in its default install folder); report its absence
plainly rather than claiming visual validation happened when it did not.
"""

# familiar's floor is Python 3.10. On anything older, stop with a plain message instead of a traceback.
import sys
if sys.version_info < (3, 10):
    sys.exit("familiar needs Python 3.10 or newer, but this is Python %d.%d (%s). Install Python 3.10+ "
             "(.familiar/scripts/doctor.py says how), then run this again with it." % (*sys.version_info[:2], sys.executable))

import argparse, glob, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path


def _find_soffice() -> str | None:
    """LibreOffice binary: PATH first, then the default install locations, which are not on PATH."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    candidates = ["/Applications/LibreOffice.app/Contents/MacOS/soffice",
                  "/usr/lib/libreoffice/program/soffice"]
    candidates += sorted(glob.glob("/opt/libreoffice*/program/soffice"))
    for var in ("ProgramFiles", "ProgramFiles(x86)"):
        base = os.environ.get(var)
        if base:
            # soffice.com is the console build; soffice.exe waits on a console and hangs headless.
            candidates += [os.path.join(base, "LibreOffice", "program", "soffice.com"),
                           os.path.join(base, "LibreOffice", "program", "soffice.exe")]
    return next((c for c in candidates if os.path.isfile(c)), None)

# python-pptx's namespace registry (pptx.oxml.ns.qn) does not include p14 -- the PowerPoint
# 2010 extension namespace that carries the (non-standard, MS-only) section list -- so it is
# addressed by its literal URI wherever section elements are read or removed.
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
P14_SECTION_LST = f"{{{P14_NS}}}sectionLst"

MAX_CHARS_PER_BULLET = 140
MAX_LINES_PER_BLOCK = 7

# Editor-facing note colors, keyed by the "note" value on a SPEC.json run -- see module
# docstring and deck-rules.md -> Editor notes. Same values and same font-color-plus-italic
# mechanism as build_docx.py's NOTE_COLOR, deliberately: a note should read the same way
# wherever a familiar-built document or deck puts one.
NOTE_COLOR = {
    "reference": (0x0B, 0x63, 0x8A),  # teal -- citation or locator
    "primary-reference": (0x6B, 0x3F, 0xA0),  # purple -- source-native marker a reference traces to
    "certainty": (0xB7, 0x79, 0x1F),  # amber -- [verify] or another confidence flag
}


def _validate_runs_spec(runs_spec) -> str | None:
    """Return an error string if `runs_spec` is malformed or names an unknown note type, else
    None. Mirrors build_docx.py's helper of the same name."""
    if not isinstance(runs_spec, list) or not runs_spec:
        return "'runs' must be a non-empty list"
    for run_spec in runs_spec:
        if not isinstance(run_spec, dict):
            return f"each run must be an object with 'text': {run_spec!r}"
        note = run_spec.get("note")
        if note is not None and note not in NOTE_COLOR:
            return f"unknown note type {note!r}; expected one of {sorted(NOTE_COLOR)}"
    return None


def _add_runs(paragraph, content) -> None:
    """Add `content` to a pptx paragraph as one or more runs -- mirrors build_docx.py's helper
    of the same name. `content` is either a plain string (single run, no formatting) or a
    pre-validated list of run specs `{"text": "...", "note": "reference"|"primary-reference"|"certainty"}`."""
    if isinstance(content, str):
        if content:
            paragraph.add_run().text = content
        return
    from pptx.dml.color import RGBColor

    for run_spec in content:
        run = paragraph.add_run()
        run.text = run_spec.get("text", "")
        note = run_spec.get("note")
        if note is not None:
            run.font.italic = True
            run.font.color.rgb = RGBColor(*NOTE_COLOR[note])


def _load_presentation(path: Path):
    """Open `path` as a python-pptx Presentation, accepting a .potx template (or .ppsx
    slideshow) in addition to a plain .pptx.

    python-pptx's own Presentation() refuses anything but the plain-presentation content
    type, so a .potx -- the normal file extension for a distributed master template, and how
    this project's own 1-Inputs/Templates/ stores them -- fails to open at all. The package
    structure is otherwise identical; retarget the declared content type in-memory rather than
    requiring a prior LibreOffice round-trip just to read the file."""
    from pptx import Presentation
    import zipfile, io

    try:
        return Presentation(str(path))
    except ValueError:
        pass

    buf = io.BytesIO()
    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.template.main+xml", b"presentationml.presentation.main+xml"
                ).replace(
                    b"presentationml.slideshow.main+xml", b"presentationml.presentation.main+xml"
                )
            dst.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def health_check() -> dict:
    try:
        import pptx  # noqa: F401
        pptx_ok, pptx_reason = True, None
    except ImportError as e:
        pptx_ok, pptx_reason = False, f"python-pptx not importable: {e}"
    try:
        import pymupdf  # noqa: F401
        pymupdf_ok = True
    except ImportError:
        pymupdf_ok = False
    renderer = _find_soffice()
    return {
        "build": pptx_ok,
        "build_reason": pptx_reason,
        "renderer": renderer,
        "visual_validation": renderer is not None and pymupdf_ok,
    }


def _all_layouts(prs):
    """Every slide layout across every slide master, in a stable flat order.

    `Presentation.slide_layouts` only exposes the first master's layouts. A template with
    more than one master (e.g. one master holding just cover layouts, another holding every
    content layout) is otherwise invisible past that first master to both --profile-template
    and layout resolution -- silently, with no error, just a short list. Flattening here is
    what makes profiling and building see the whole template."""
    layouts = []
    for master in prs.slide_masters:
        layouts.extend(master.slide_layouts)
    return layouts


def _placeholder_conflicts(layout):
    """Group this layout's placeholders by (idx, type) and return only groups with more than
    one member -- keyed by that (idx, type) pair, valued as the list of colliding placeholders.

    python-pptx (and PowerPoint itself) correlates a slide's inherited placeholder to its
    layout counterpart by (type, idx) alone. A layout authored from retrofitted text boxes
    can end up with two placeholders of the same type both left at idx 0; when that happens
    the correlation is genuinely ambiguous and which placeholder a slide's content actually
    binds to is not guaranteed -- confirmed against a real template where this put bullet
    content into a half-inch subhead band instead of the main content box."""
    from collections import defaultdict

    groups = defaultdict(list)
    for ph in layout.placeholders:
        groups[(ph.placeholder_format.idx, str(ph.placeholder_format.type))].append(ph)
    return {key: phs for key, phs in groups.items() if len(phs) > 1}


def _describe_conflicts(conflicts) -> str:
    return "; ".join(
        f"idx={idx} type={ptype} shared by {[ph.name for ph in phs]}"
        for (idx, ptype), phs in conflicts.items()
    )


def _resolve_layout(prs, layout_ref, default_index):
    """layout_ref may be an int index, a layout name (str), or absent (-> default_index).
    Indices and names are resolved against _all_layouts, i.e. every master's layouts flattened
    -- the same order --profile-template reports, so its `index` values line up directly."""
    layouts = _all_layouts(prs)
    if layout_ref is None:
        idx = default_index
        if idx >= len(layouts):
            raise ValueError(f"template has only {len(layouts)} layout(s); no default index {idx}")
        layout = layouts[idx]
    elif isinstance(layout_ref, int):
        if not (0 <= layout_ref < len(layouts)):
            raise ValueError(f"layout index {layout_ref} out of range (template has {len(layouts)})")
        layout = layouts[layout_ref]
    else:
        layout = None
        for candidate in layouts:
            if candidate.name == layout_ref:
                layout = candidate
                break
        if layout is None:
            names = [l.name for l in layouts]
            raise ValueError(f"no layout named {layout_ref!r} in template; available: {names}")

    conflicts = _placeholder_conflicts(layout)
    if conflicts:
        raise ValueError(
            f"layout {layout.name!r} has ambiguous placeholders and cannot be built on safely: "
            f"{_describe_conflicts(conflicts)} -- python-pptx and PowerPoint cannot tell these "
            f"apart by (type, idx) alone, so content may bind to the wrong one. Fix the template "
            f"in PowerPoint's Slide Master view (give each placeholder a distinct idx) or choose "
            f"a different layout -- see --profile-template's \"usable\" field."
        )
    return layout


def _require_new_output(out_path: Path, *sources: Path) -> None:
    output = Path(out_path).resolve()
    if any(output == Path(source).resolve() for source in sources):
        raise ValueError("--output must not be the same file as a template")
    if output.exists():
        raise ValueError(f"--output already exists: {out_path}; choose a new versioned filename")


def _clear_template_slides(prs) -> None:
    """Keep the template's design parts while removing its project-specific slides and the
    section list (p14:sectionLst, a PowerPoint extension) that groups them -- leftover
    sections still point at the removed slide IDs and can mis-sort newly added slides into
    stale template sections that no longer show any content in the Slides pane."""
    for slide_id in list(prs.slides._sldIdLst):
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)
    from pptx.oxml.ns import qn

    ext_lst = prs._element.find(qn("p:extLst"))
    if ext_lst is not None:
        for ext in list(ext_lst):
            if ext.find(P14_SECTION_LST) is not None:
                ext_lst.remove(ext)
        if len(ext_lst) == 0:
            prs._element.remove(ext_lst)


def _set_title(slide, text, layout_name: str) -> None:
    if not text:
        return
    title = slide.shapes.title
    if title is None:
        raise ValueError(f"layout {layout_name!r} has no title placeholder; choose a title-capable layout")
    title.text = text


def _body_placeholder(slide, layout_name: str):
    title = slide.shapes.title
    title_id = title.shape_id if title is not None else None
    for placeholder in slide.placeholders:
        if placeholder.shape_id != title_id and placeholder.has_text_frame:
            return placeholder
    raise ValueError(f"layout {layout_name!r} has no non-title text placeholder for bullets")


def profile_template(path: Path) -> dict:
    """The layout-profile entry point deck-rules.md requires before building against an
    unfamiliar template: every layout's index, name, dimensions, and each placeholder's
    idx/type/name -- what a spec's `layout`/`title_layout` fields, and a hand-authored slide
    spec generally, need instead of guessing from visual similarity.

    Covers every slide master's layouts, flattened via _all_layouts, in the same order
    _resolve_layout indexes against. A layout with `"usable": false` has two or more
    placeholders sharing both idx and type -- see _placeholder_conflicts -- and must not be
    passed as a spec's `layout`/`title_layout`; _resolve_layout refuses it for the same reason."""
    prs = _load_presentation(path)
    layouts = []
    master_index_of = {id(master): mi for mi, master in enumerate(prs.slide_masters)}
    for i, layout in enumerate(_all_layouts(prs)):
        conflicts = _placeholder_conflicts(layout)
        layouts.append({
            "index": i,
            "master_index": master_index_of[id(layout.slide_master)],
            "name": layout.name,
            "usable": not conflicts,
            "conflicts": [
                {"idx": idx, "type": ptype, "placeholder_names": [ph.name for ph in phs]}
                for (idx, ptype), phs in conflicts.items()
            ],
            "placeholders": [
                {"idx": ph.placeholder_format.idx,
                 "type": str(ph.placeholder_format.type),
                 "name": ph.name}
                for ph in layout.placeholders
            ],
        })
    return {
        "slide_width": prs.slide_width,
        "slide_height": prs.slide_height,
        "layouts": layouts,
    }


def build(spec: dict, out_path: Path, template_path: Path | None = None) -> None:
    from pptx import Presentation

    _require_new_output(out_path, *([template_path] if template_path else []))
    prs = _load_presentation(template_path) if template_path else Presentation()
    if template_path:
        _clear_template_slides(prs)

    title_layout = _resolve_layout(prs, spec.get("title_layout"), 0)
    title_slide = prs.slides.add_slide(title_layout)
    _set_title(title_slide, spec.get("title", ""), title_layout.name)

    for slide_spec in spec.get("slides", []):
        layout = _resolve_layout(prs, slide_spec.get("layout"), 1)
        slide = prs.slides.add_slide(layout)
        _set_title(slide, slide_spec.get("title", ""), layout.name)
        bullets = slide_spec.get("bullets", [])
        if not bullets:
            continue
        body = _body_placeholder(slide, layout.name).text_frame
        for i, bullet in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            if isinstance(bullet, dict):
                runs_spec = bullet.get("runs", [])
                err = _validate_runs_spec(runs_spec)
                if err:
                    raise ValueError(f"invalid bullet runs: {err}")
                _add_runs(p, runs_spec)
            else:
                _add_runs(p, bullet)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))


def validate(path: Path, template_path: Path | None = None) -> list[str]:
    from pptx.oxml.ns import qn

    findings = []
    prs = _load_presentation(path)
    if len(prs.slides) < 2:
        findings.append(f"only {len(prs.slides)} slide(s); expected a title slide plus content")
    live_slide_ids = {slide_id.get("id") for slide_id in prs.slides._sldIdLst}
    ext_lst = prs._element.find(qn("p:extLst"))
    section_lst = None
    if ext_lst is not None:
        for ext in ext_lst:
            section_lst = ext.find(P14_SECTION_LST)
            if section_lst is not None:
                break
    if section_lst is not None:
        for section in section_lst:
            section_name = section.get("name", "(unnamed)")
            sld_id_lst = section.find(f"{{{P14_NS}}}sldIdLst")
            entry_ids = [sld_id.get("id") for sld_id in sld_id_lst] if sld_id_lst is not None else []
            if not entry_ids:
                findings.append(f"section {section_name!r}: empty -- leftover template section not removed")
                continue
            for entry_id in entry_ids:
                if entry_id not in live_slide_ids:
                    findings.append(
                        f"section {section_name!r}: references slide id {entry_id} which no longer "
                        f"exists in this deck -- leftover template section, not cleaned up on build"
                    )
    if template_path:
        tpl = _load_presentation(template_path)
        if (prs.slide_width, prs.slide_height) != (tpl.slide_width, tpl.slide_height):
            findings.append(
                f"dimensions {prs.slide_width}x{prs.slide_height} do not match template "
                f"{tpl.slide_width}x{tpl.slide_height} -- the template was not actually used as the base"
            )
    # Density is judged per text block (placeholder, card, tile, grouped shape), not per slide: a
    # visual slide splits its text across several short blocks, which a whole-slide line count
    # flags as dense when it is the opposite.
    for i, slide in enumerate(prs.slides):
        for shape in _text_shapes(slide.shapes):
            lines = 0
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if not text.strip():
                    continue
                lines += 1
                if len(text) > MAX_CHARS_PER_BULLET:
                    findings.append(f"slide {i + 1}: line has {len(text)} chars (possible overflow, threshold {MAX_CHARS_PER_BULLET}): {text[:60]}...")
            if lines > MAX_LINES_PER_BLOCK:
                findings.append(f"slide {i + 1}: text block {shape.name!r} has {lines} lines (possible density overflow, threshold {MAX_LINES_PER_BLOCK})")
    return findings


def _text_shapes(shapes):
    """Yield every shape with a text frame, descending into groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _text_shapes(shape.shapes)
        elif shape.has_text_frame:
            yield shape


def render(path: Path, out_dir: Path) -> list[Path]:
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError("no soffice/libreoffice binary found (PATH or default install folder) -- visual rendering unavailable on this host")
    import pymupdf

    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
            check=True, capture_output=True,
        )
        pdf_path = Path(tmp) / (path.stem + ".pdf")
        pages = []
        with pymupdf.open(str(pdf_path)) as pdf:
            for i, page in enumerate(pdf):
                png_path = out_dir / f"{path.stem}-slide{i + 1:02d}.png"
                page.get_pixmap(dpi=150).save(str(png_path))
                pages.append(png_path)
    return pages


def main():
    ap = argparse.ArgumentParser()
    group = ap.add_mutually_exclusive_group(required=True)
    group.add_argument("--health-check", action="store_true")
    group.add_argument("--build", nargs=2, metavar=("SPEC_JSON", "OUT_PPTX"))
    group.add_argument("--profile-template", metavar="TEMPLATE_PPTX")
    group.add_argument("--validate", metavar="PPTX")
    group.add_argument("--render", nargs=2, metavar=("PPTX", "OUT_DIR"))
    ap.add_argument("--template", metavar="TEMPLATE_PPTX",
                     help="with --build: build starting from this presentation's master/layouts "
                          "instead of a blank default. With --validate: also check the built "
                          "deck kept this template's slide dimensions.")
    args = ap.parse_args()

    if args.health_check:
        print(json.dumps(health_check(), indent=2))
        return

    if args.profile_template:
        print(json.dumps(profile_template(Path(args.profile_template)), indent=2))
        return

    if args.build:
        spec_path, out_path = args.build
        spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        try:
            build(spec, Path(out_path), Path(args.template) if args.template else None)
        except (ValueError, AttributeError) as e:
            print(f"ERROR: {e}", file=sys.stderr)
            sys.exit(4)
        print(f"-> {out_path}")
        return

    if args.validate:
        findings = validate(Path(args.validate), Path(args.template) if args.template else None)
        if findings:
            print(f"STRUCTURAL ISSUES ({len(findings)}):")
            for f in findings:
                print(f"  - {f}")
            sys.exit(1)
        print("structural validation: no issues found")
        return

    if args.render:
        pptx_path, out_dir = args.render
        try:
            pages = render(Path(pptx_path), Path(out_dir))
        except (RuntimeError, subprocess.CalledProcessError) as e:
            print(f"ERROR: {e}", file=sys.stderr)
            sys.exit(2)
        for p in pages:
            print(p)


if __name__ == "__main__":
    main()
