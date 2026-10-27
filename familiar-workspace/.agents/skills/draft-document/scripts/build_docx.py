#!/usr/bin/env python3
"""build_docx.py -- host-neutral Word document build/validate/render for the document skill.

Uses python-docx (structural build + reopen-and-check) and, when present, a LibreOffice
`soffice` binary (render to PDF for page-by-page visual inspection). Neither is guaranteed to
be installed on every host -- run --health-check first and report the exact missing piece
rather than attempting an unvalidated build.

    python build_docx.py --health-check
    python build_docx.py --build SPEC.json OUT.docx
    python build_docx.py --from-template TPL.docx SPEC.json OUT.docx --mode shell|extend
    python build_docx.py --revise SOURCE.docx EDITS.json OUT.docx
    python build_docx.py --validate OUT.docx [--source SOURCE_OR_TEMPLATE.docx]
    python build_docx.py --render OUT.docx OUT_DIR

SPEC.json shape (--build, --from-template):
    {"title": "Doc title", "sections": [{"heading": "...", "paragraphs": ["...", "..."]}, ...]}

A paragraph entry is normally a plain string. It may instead be `{"runs": [...]}` to mix body
text with an editor-facing note -- a source citation/locator or a `voice.md` -> Uncertainty
marker `[verify]` flag -- so the note is visually distinct (colored, italic text) instead of
disappearing into the body font as plain bracket text on a page that gets read and edited
several times before the note is actually resolved:
    {"runs": [{"text": "Effect size was 4.2 points "},
               {"text": "[Smith 2024, p.12]", "note": "reference"}]}
Each run is `{"text": "..."}`, optionally with `"note"`: `"reference"` (teal) for a citation or
locator, `"primary-reference"` (purple) for the source-native marker a `reference` claim traces
back to (a knowledge page's `ref_carry_through` -- the underlying study, not the compiled document
citing it), or `"certainty"` (amber, the
same color an extension's own VERIFY: callout may use) for `[verify]` or another confidence flag. A claim with both layers of provenance
gets both runs, back to back -- see docx-rules.md -> Editor notes -> Dual citation for the exact
shape and when it applies. Deliberately not a highlight: a background fill reads as "look here
right now," while a note that will sit on the page through several editing passes reads better as
a quieter, still-unmistakable color/italic combination. The marker text itself is unchanged by
this -- only its formatting is. An unrecognized note type is refused, not guessed at. See
docx-rules.md -> Editor notes.

--from-template starts from an existing document instead of a blank one, in one of two modes
(no default -- a silent guess is how a report ships with the previous project's table still in
it, per docx-rules.md):
  shell   clears the template's body content (paragraphs, tables) but keeps its section
          properties (page setup, header/footer references) and styles, then appends SPEC's
          content -- "use this document's look for a new report."
  extend  keeps the template's body content and appends SPEC's content after it --
          "use this document as a starting point." The template's own tables/boilerplate come
          along; wrong if a blank shell was actually wanted.
The template file itself is never modified; the output must be a new file. Every build mode
refuses an existing output path before writing, so a dated draft version cannot be overwritten
by a rerun.

--revise applies targeted edits to an existing document via EDITS.json:
    {"edits": [{"anchor": "exact paragraph text to locate",
                "occurrence": 1,                 # optional, 1-based; required only when the
                                                  # anchor text isn't unique in the document
                "action": "replace" | "insert_after",
                "text": "new text"}]}
An edit may carry `"runs": [...]` instead of `"text"` -- the same run-list shape as a SPEC.json
paragraph (see above) -- to give the new or replacement content an editor-facing `reference` or
`certainty` note run instead of plain text.
`anchor` is matched against a whole paragraph's text (not a substring) in true document order,
including paragraphs inside table cells. An anchor matching zero or more-than-one location without
`occurrence` disambiguating it is refused rather than guessed at. `replace` on a paragraph
containing non-text run content (e.g. an image) is refused rather than silently destroying it --
narrow the edit instead. Whole-cell replacement is intentionally unsupported because it can
destroy multi-paragraph structure; target the exact paragraph instead.

Structural --validate is always available wherever --build is (same python-docx dependency).
It always checks editor-note formatting -- a bracketed run (a citation/locator, `[verify]`, or
`[action: ...]`) left with no explicit color and no italic is reported, since that check is
meaningful on a single file and needs no --source. A run colored for a reason this script
doesn't recognize (e.g. a template's own instructional placeholders) is left alone -- the test
is "has no explicit color," not "isn't teal or amber."
With --source, --validate additionally checks the output kept the source's headers, footers,
table count, embedded-object count, and character formatting (the mix of font/size in use,
compared as a share of total text) -- catching exactly the class of loss `--from-template extend`
and `--revise` must not cause. Counting objects alone is not enough: a revision can keep every
table and image and still drop the font the document is required to be written in.
--render additionally needs a `soffice` (LibreOffice) binary (on PATH or in its default install folder); report its absence
plainly rather than claiming visual validation happened when it did not.
"""

# familiar's floor is Python 3.10. On anything older, stop with a plain message instead of a traceback.
import sys
if sys.version_info < (3, 10):
    sys.exit("familiar needs Python 3.10 or newer, but this is Python %d.%d (%s). Install Python 3.10+ "
             "(.familiar/scripts/doctor.py says how), then run this again with it." % (*sys.version_info[:2], sys.executable))

import argparse, glob, json, os, re, shutil, subprocess, sys
from pathlib import Path


def _find_soffice() -> str | None:
    """LibreOffice binary: PATH first, then the default install locations, which are not on PATH."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    candidates = ["/Applications/LibreOffice.app/Contents/MacOS/soffice",
                  "/usr/lib/libreoffice/program/soffice"]
    candidates += sorted(glob.glob("/opt/libreoffice*/program/soffice"))
    for var in ("ProgramFiles", "ProgramFiles(x86)"):
        base = os.environ.get(var)
        if base:
            # soffice.com is the console build; soffice.exe waits on a console and hangs headless.
            candidates += [os.path.join(base, "LibreOffice", "program", "soffice.com"),
                           os.path.join(base, "LibreOffice", "program", "soffice.exe")]
    return next((c for c in candidates if os.path.isfile(c)), None)

# Editor-facing note colors, keyed by the "note" value on a SPEC.json/EDITS.json run -- see
# module docstring and docx-rules.md -> Editor notes. Font color + italic, not a highlight: a
# note sits on the page through several editing passes before it's resolved, so it should read
# as visibly distinct without competing for attention like a "look here now" background fill.
# "certainty" is the shared VERIFY color, so a [verify]-style flag looks the
# same everywhere in a familiar-built document, not just in this skill.
NOTE_COLOR = {
    "reference": (0x0B, 0x63, 0x8A),  # teal -- citation or locator
    "primary-reference": (0x6B, 0x3F, 0xA0),  # purple -- source-native marker a reference traces to
    "certainty": (0xB7, 0x79, 0x1F),  # amber -- [verify] or another confidence flag
}


def health_check() -> dict:
    try:
        import docx  # noqa: F401
        docx_ok, docx_reason = True, None
    except ImportError as e:
        docx_ok, docx_reason = False, f"python-docx not importable: {e}"
    renderer = _find_soffice()
    return {
        "build": docx_ok,
        "build_reason": docx_reason,
        "renderer": renderer,
        "visual_validation": renderer is not None,
    }


def _validate_runs_spec(runs_spec) -> str | None:
    """Return an error string if `runs_spec` is malformed or names an unknown note type, else
    None. Checked before any document mutation so a bad note never leaves a half-applied edit."""
    if not isinstance(runs_spec, list) or not runs_spec:
        return "'runs' must be a non-empty list"
    for run_spec in runs_spec:
        if not isinstance(run_spec, dict):
            return f"each run must be an object with 'text': {run_spec!r}"
        note = run_spec.get("note")
        if note is not None and note not in NOTE_COLOR:
            return f"unknown note type {note!r}; expected one of {sorted(NOTE_COLOR)}"
    return None


# Character formatting carried from a replaced run onto its replacements. A rebuilt run starts
# with an empty rPr, so without this it silently falls back to the paragraph style's font --
# which is how a --revise pass once stripped Arial 10pt out of a template that requires it and
# repaginated the document. Color and italic are deliberately NOT carried: a note run sets its
# own, and body text must not inherit the previous run's note coloring.
_CARRIED_RPR_TAGS = ("w:rFonts", "w:sz", "w:szCs", "w:lang", "w:b", "w:bCs", "w:caps", "w:smallCaps")


def _base_rpr(paragraph):
    """The first run-properties element in `paragraph` that carries character formatting worth
    inheriting, or None. Used so replacement runs keep the font they are replacing."""
    from docx.oxml.ns import qn

    for run in paragraph.runs:
        rpr = run._element.find(qn("w:rPr"))
        if rpr is not None and any(rpr.find(qn(tag)) is not None for tag in _CARRIED_RPR_TAGS):
            return rpr
    return None


def _apply_base_rpr(run, base) -> None:
    """Give `run` the carried character formatting from `base`, without disturbing anything the
    run already sets for itself (a note's color and italic are applied after this)."""
    import copy

    from docx.oxml.ns import qn

    if base is None:
        return
    rpr = run._element.get_or_add_rPr()
    for tag in _CARRIED_RPR_TAGS:
        source = base.find(qn(tag))
        if source is not None and rpr.find(qn(tag)) is None:
            rpr.insert(0, copy.deepcopy(source))


def _add_runs(paragraph, content, base_rpr=None) -> None:
    """Add `content` to `paragraph` as one or more runs. `content` is either a plain string
    (single run, no formatting -- matches `Document.add_paragraph(text)`'s own behavior of
    adding nothing for an empty string) or a pre-validated list of run specs
    `{"text": "...", "note": "reference"|"primary-reference"|"certainty"}`. A `note` run is colored and italicized so
    an editor-facing note reads as visibly distinct from body text -- see docx-rules.md ->
    Editor notes.

    `base_rpr` is the run-properties element the new runs inherit character formatting from --
    passed by --revise so a rebuilt paragraph keeps the font of the one it replaces. Omitted for
    --build/--from-template, where new content legitimately takes the paragraph style's font."""
    if isinstance(content, str):
        if content:
            _apply_base_rpr(paragraph.add_run(content), base_rpr)
        return
    from docx.shared import RGBColor

    for run_spec in content:
        run = paragraph.add_run(run_spec.get("text", ""))
        _apply_base_rpr(run, base_rpr)
        note = run_spec.get("note")
        if note is not None:
            run.italic = True
            run.font.color.rgb = RGBColor(*NOTE_COLOR[note])


def _append_spec_content(doc, spec: dict) -> None:
    """Append SPEC's headings/paragraphs to `doc` (Document or a template already opened),
    shared by --build, --from-template shell, and --from-template extend -- the same content
    goes in regardless of what, if anything, was already in the document."""
    if spec.get("title"):
        doc.add_heading(spec["title"], level=0)
    for section in spec.get("sections", []):
        if section.get("heading"):
            doc.add_heading(section["heading"], level=1)
        for para in section.get("paragraphs", []):
            p = doc.add_paragraph()
            if isinstance(para, dict):
                runs_spec = para.get("runs", [])
                err = _validate_runs_spec(runs_spec)
                if err:
                    raise ValueError(f"invalid paragraph runs: {err}")
                _add_runs(p, runs_spec)
            else:
                _add_runs(p, para)


def _require_new_output(out_path: Path, *sources: Path) -> None:
    """Refuse to replace a source or an existing deliverable.

    familiar's copy-on-edit rule needs an executable check here: a builder cannot know whether
    an existing file is a source, an approved draft, or another user's work.
    """
    output = Path(out_path).resolve()
    if any(output == Path(source).resolve() for source in sources):
        raise ValueError("--output must not be the same file as a source or template")
    if output.exists():
        raise ValueError(f"--output already exists: {out_path}; choose a new versioned filename")


def _partial_output(out_path: Path, *sources: Path) -> Path:
    """Where an incomplete result goes: `<stem>.partial<ext>`, or the first free
    `<stem>.partial-N<ext>` when that name is taken or is itself one of the sources. Retrying from
    an earlier partial (source `x.partial.docx`, output `x.docx`) must not write the new partial
    over the file being revised, and a second failed run must not replace the first one's partial."""
    out_path = Path(out_path)
    resolved_sources = {Path(source).resolve() for source in sources}
    n = 1
    while True:
        suffix = ".partial" if n == 1 else f".partial-{n}"
        candidate = out_path.with_name(out_path.stem + suffix + out_path.suffix)
        if not candidate.exists() and candidate.resolve() not in resolved_sources:
            return candidate
        n += 1


def build(spec: dict, out_path: Path) -> None:
    from docx import Document

    _require_new_output(out_path)
    doc = Document()
    _append_spec_content(doc, spec)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def build_from_template(template_path: Path, spec: dict, out_path: Path, mode: str) -> None:
    from docx import Document

    _require_new_output(out_path, template_path)
    doc = Document(str(template_path))
    if mode == "shell":
        # Clear body content but keep sectPr (page setup, header/footer references) and styles --
        # "use this document's look for a new report."
        body = doc.element.body
        for child in list(body):
            if child.tag.endswith("}sectPr"):
                continue
            body.remove(child)
    # "extend" keeps the template's existing body content as-is.
    _append_spec_content(doc, spec)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


# --- paragraph walk and run splitting. Assigning to Run.text clears every non-rPr child of
# <w:r>, silently destroying a <w:drawing> or other non-text content in a run that needs
# splitting at a character offset. -----------------------------------------------------------
def _iter_paragraphs(container):
    """Paragraphs in true document order, descending into tables (and nested tables) the same
    way, visiting a merged cell once rather than once per spanned grid position.
    Yields (container, paragraph) so a caller can add a sibling paragraph in the right place.

    Cells are walked from the table's own XML (<w:tr> -> <w:tc>), where a merged cell is stored
    exactly once, instead of through python-docx's row.cells, which repeats the same cell for
    every grid position it spans. The previous version deduplicated row.cells by id(cell._tc),
    which is not safe: lxml creates element proxies on demand and frees them when unreferenced,
    so CPython reuses their ids and later, unrelated cells were silently skipped as duplicates.
    That made this iterator non-deterministic -- the same file yielded 179 paragraphs on one run
    and 154 on the next -- and since --revise locates its anchors through it, an anchor inside a
    table could be reported as unmatched, or an `occurrence` index could land on the wrong
    paragraph. Walking the XML needs no deduplication at all."""
    from docx.oxml.ns import qn
    from docx.table import _Cell
    from docx.text.paragraph import Paragraph

    for item in container.iter_inner_content():
        if isinstance(item, Paragraph):
            yield container, item
        else:  # Table
            for tr in item._tbl.iterchildren(qn("w:tr")):
                for tc in tr.iterchildren(qn("w:tc")):
                    yield from _iter_paragraphs(_Cell(tc, item))


def _run_has_only_text_content(run) -> bool:
    """True if this run's only content children are text-like (safe to collapse to one run
    without losing anything); False if it carries a drawing or other real non-text content.

    <w:lastRenderedPageBreak> counts as text-like: it is a pagination hint Word writes on save
    and regenerates on the next one, not content. Refusing a paragraph over it blocked edits to
    exactly the long paragraphs most likely to need them."""
    from docx.oxml.ns import qn

    text_like = {
        qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr"), qn("w:noBreakHyphen"), qn("w:ptab"),
        qn("w:lastRenderedPageBreak"),
    }
    rpr = qn("w:rPr")
    return all(child.tag in text_like or child.tag == rpr for child in run._element)


def _nested_text_container(paragraph) -> str | None:
    """Name of the first element holding visible text *outside* the paragraph's direct runs, or
    None. `paragraph.runs` lists only direct `w:r` children, so text inside a hyperlink, simple
    field, smart tag, content control, or tracked insertion survives a clear-and-rewrite of those
    runs -- "Old <link>linked content</link>" came back as "Replacementlinked content" while the
    edit reported success and validation found nothing."""
    from docx.oxml.ns import qn

    p = paragraph._p
    for t in p.iter(qn("w:t"), qn("w:delText")):
        container = t.getparent()
        while container is not None and container.getparent() is not p:
            container = container.getparent()
        if container is not None and container.tag != qn("w:r"):
            return container.tag.split("}")[-1]
    return None


def revise(source_path: Path, edits: list[dict], out_path: Path,
           partial_path: Path | None = None) -> tuple[int, list[tuple[int, str, str]]]:
    """Apply targeted edits located by anchor text. Returns (applied_count, failures) where each
    failure is (edit_index, anchor, reason). Refuses ambiguous/missing anchors and any `replace`
    on a paragraph carrying non-text run content rather than guessing or silently destroying it.
    Saves to `out_path` only if every edit applied; on any failure, saves a `.partial` sibling
    instead and leaves `out_path` unwritten -- same contract as annotate.py's placement pass, so
    a caller checking only for the requested file's existence can never mistake a partial result
    for the deliverable. The partial name comes from `_partial_output` unless the caller already
    chose one (the CLI does, so it can report the name it will use)."""
    from docx import Document

    _require_new_output(out_path, source_path)

    doc = Document(str(source_path))

    # Build the full candidate list ONCE per revise() call. Re-locating fresh per edit would let
    # an earlier edit's own text changes shift what a later edit's anchor happens to match.
    # A one-paragraph cell is represented by its paragraph only: adding its whole-cell text as a
    # second candidate turns one visible anchor into a false ambiguity.
    candidates = []
    cells = {}
    for container, p in _iter_paragraphs(doc):
        candidates.append(("paragraph", container, p, p.text.strip()))
        if hasattr(container, "_tc"):
            cells[id(container._tc)] = container
    for cell in cells.values():
        cell_text = cell.text.strip()
        if cell_text and not any(p.text.strip() == cell_text for p in cell.paragraphs):
            candidates.append(("cell", None, cell, cell_text))

    applied, failures = 0, []
    for n, edit in enumerate(edits, 1):
        anchor = str(edit.get("anchor", "")).strip()
        action = edit.get("action")
        text = edit.get("text", "")
        runs_spec = edit.get("runs")
        if not anchor:
            failures.append((n, anchor, "missing or blank required 'anchor'"))
            continue
        if action not in ("replace", "insert_after"):
            failures.append((n, anchor, f"unknown action '{action}'"))
            continue
        if runs_spec is not None:
            err = _validate_runs_spec(runs_spec)
            if err:
                failures.append((n, anchor, f"invalid 'runs': {err}"))
                continue
        matches = [c for c in candidates if c[3] == anchor]
        if not matches:
            failures.append((n, anchor, "anchor not found"))
            continue
        occurrence = edit.get("occurrence")
        if occurrence is None:
            if len(matches) > 1:
                failures.append((n, anchor, f"anchor matches {len(matches)} locations; "
                                             f"add 'occurrence' to disambiguate"))
                continue
            match = matches[0]
        else:
            idx = int(occurrence) - 1
            if not (0 <= idx < len(matches)):
                failures.append((n, anchor, f"occurrence {occurrence} out of range "
                                             f"({len(matches)} match(es))"))
                continue
            match = matches[idx]
        kind, container, target, _ = match

        if action == "insert_after":
            if kind != "paragraph":
                failures.append((n, anchor, "'insert_after' applies to a paragraph anchor, not a table cell"))
                continue
            new_p = container.add_paragraph(style=target.style)
            _add_runs(new_p, runs_spec if runs_spec is not None else text, _base_rpr(target))
            target._p.addnext(new_p._p)
            applied += 1
            continue

        # action == "replace"
        if kind == "cell":
            failures.append((n, anchor, "whole-cell replacement is not supported because it can "
                                         "remove cell structure; target one paragraph instead"))
            continue
        if not all(_run_has_only_text_content(r) for r in target.runs):
            failures.append((n, anchor, "paragraph contains non-text content (e.g. an image); "
                                         "targeted replace of its full text is not supported"))
            continue
        nested = _nested_text_container(target)
        if nested:
            failures.append((n, anchor, f"paragraph has text inside a w:{nested} element (e.g. a "
                                         f"hyperlink or tracked change); replacing its full text would "
                                         f"leave that text behind, so it is not supported"))
            continue
        if runs_spec is not None:
            # A note run needs its own color and italic, not the anchor run's -- so clear and
            # rebuild rather than reusing runs[0] the way the plain-text path does below. Capture
            # the anchor's character formatting first and hand it to the rebuilt runs: clearing
            # must drop the note styling, not the font the paragraph was written in.
            base = _base_rpr(target)
            for r in list(target.runs):
                r._element.getparent().remove(r._element)
            _add_runs(target, runs_spec, base)
        else:
            for r in target.runs[1:]:
                r._element.getparent().remove(r._element)
            if target.runs:
                target.runs[0].text = text
            else:
                target.add_run(text)
        applied += 1

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if failures:
        partial_path = Path(partial_path) if partial_path else _partial_output(out_path, source_path)
        if partial_path.exists() or partial_path.resolve() == Path(source_path).resolve():
            raise ValueError(f"partial output {partial_path} would replace an existing file or the source")
        doc.save(str(partial_path))
    else:
        doc.save(str(out_path))
    return applied, failures


def _font_usage(doc) -> dict:
    """Characters of body text per (font name, size) actually in use, across paragraphs and
    table cells. `None` means the run sets nothing and inherits its style, which is a legitimate
    state -- what matters is that the mix does not shift between a source and its revision."""
    usage: dict = {}
    for _container, paragraph in _iter_paragraphs(doc):
        for run in paragraph.runs:
            size = run.font.size.pt if run.font.size is not None else None
            key = (run.font.name, size)
            usage[key] = usage.get(key, 0) + len(run.text)
    return usage


def _font_drift(source_doc, out_doc) -> list:
    """Findings for character formatting the output lost relative to its source.

    Counting tables and images is not enough. Rebuilding a run gives it an empty rPr, so a
    revision can keep every table, image, header and footer and still quietly drop the font the
    document is required to be written in -- which is what happened to a regulator's template that
    mandates 10-point Arial. Compared as a share of total text, not an absolute count, because a
    revision legitimately changes length. Only combinations carrying at least 1% of the source
    are checked, so losing an incidental bold word is not a finding."""
    src, out = _font_usage(source_doc), _font_usage(out_doc)
    src_total, out_total = sum(src.values()), sum(out.values())
    if not src_total or not out_total:
        return []
    findings = []
    for key, chars in sorted(src.items(), key=lambda kv: -kv[1]):
        src_share = chars / src_total
        if src_share < 0.01:
            continue
        out_share = out.get(key, 0) / out_total
        if out_share < src_share * 0.9:
            name, size = key
            described = f"{name or 'the style default font'}" + (f" {size}pt" if size else "")
            findings.append(
                f"source had {src_share:.0%} of its text in {described}, output has "
                f"{out_share:.0%} -- character formatting appears lost, not just changed"
            )
    return findings


_BRACKET_NOTE_RE = re.compile(r"^\[.+\]$")


def _note_format_drift(doc) -> list:
    """Findings for a bracketed editor note (a citation/locator, `[verify]`, or `[action: ...]`)
    left in the body font instead of given its own colored, italic run -- see docx-rules.md ->
    Editor notes. Runs unconditionally, on every --validate call, because a plain black bracket
    in the body font is exactly the defect that convention exists to prevent, and it is visible
    from the output alone -- no --source needed.

    Tests "has no explicit color", not "isn't teal or amber": a template's own instructional
    placeholders (e.g. a regulator template's dark-red `[name]`/`[proposed]` callouts) are
    deliberately colored for a reason this script doesn't recognize, and are not ours to flag.
    An unreadable (themed/indexed) color is treated as "has a color" for the same reason -- this
    is a lint for the specific defect that shipped, not a full color audit."""
    findings = []
    for _container, paragraph in _iter_paragraphs(doc):
        for run in paragraph.runs:
            text = run.text.strip()
            if not text or not _BRACKET_NOTE_RE.match(text):
                continue
            try:
                has_color = run.font.color is not None and run.font.color.rgb is not None
            except Exception:
                has_color = True
            if not has_color and not run.italic:
                findings.append(
                    f"possible editor note left in body font: {text!r} -- give it a colored, "
                    f"italic run per docx-rules.md -> Editor notes"
                )
    return findings


def _preservation_counts(doc) -> dict:
    return {
        "tables": len(doc.tables),
        "headers with text": sum(
            1 for s in doc.sections if any(p.text.strip() for p in s.header.paragraphs)
        ),
        "footers with text": sum(
            1 for s in doc.sections if any(p.text.strip() for p in s.footer.paragraphs)
        ),
        "embedded images": len(doc.inline_shapes),
    }


def validate(path: Path, source_path: Path | None = None) -> list[str]:
    from docx import Document

    findings = []
    doc = Document(str(path))
    non_empty = [p for p in doc.paragraphs if p.text.strip()]
    if not non_empty:
        findings.append("document has no non-empty paragraphs")
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    if not headings:
        findings.append("document has no heading or title paragraphs")
    findings.extend(_note_format_drift(doc))

    if source_path:
        source_doc = Document(str(source_path))
        src_counts = _preservation_counts(source_doc)
        out_counts = _preservation_counts(doc)
        for key, src_n in src_counts.items():
            if out_counts[key] < src_n:
                findings.append(
                    f"source had {src_n} {key}, output has only {out_counts[key]} -- "
                    f"content appears lost, not just changed"
                )
        findings.extend(_font_drift(source_doc, doc))
    return findings


def render(path: Path, out_dir: Path) -> Path:
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError("no soffice/libreoffice binary found (PATH or default install folder) -- visual rendering unavailable on this host")
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
        check=True, capture_output=True,
    )
    pdf_path = out_dir / (path.stem + ".pdf")
    if not pdf_path.is_file():
        raise RuntimeError(f"soffice did not produce {pdf_path}")
    return pdf_path


def main():
    ap = argparse.ArgumentParser()
    group = ap.add_mutually_exclusive_group(required=True)
    group.add_argument("--health-check", action="store_true")
    group.add_argument("--build", nargs=2, metavar=("SPEC_JSON", "OUT_DOCX"))
    group.add_argument("--from-template", nargs=3, metavar=("TEMPLATE_DOCX", "SPEC_JSON", "OUT_DOCX"))
    group.add_argument("--revise", nargs=3, metavar=("SOURCE_DOCX", "EDITS_JSON", "OUT_DOCX"))
    group.add_argument("--validate", metavar="DOCX")
    group.add_argument("--render", nargs=2, metavar=("DOCX", "OUT_DIR"))
    ap.add_argument("--mode", choices=["shell", "extend"], help="required with --from-template")
    ap.add_argument("--source", metavar="SOURCE_OR_TEMPLATE_DOCX",
                     help="with --validate: also check the output kept this document's headers, "
                          "footers, tables, and embedded images")
    args = ap.parse_args()

    if args.health_check:
        print(json.dumps(health_check(), indent=2))
        return

    if args.build:
        spec_path, out_path = args.build
        spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        try:
            build(spec, Path(out_path))
        except ValueError as e:
            print(f"ERROR: {e}", file=sys.stderr)
            sys.exit(1)
        print(f"-> {out_path}")
        return

    if args.from_template:
        template_path, spec_path, out_path = args.from_template
        if not args.mode:
            ap.error("--from-template requires --mode shell|extend -- there is no default; "
                     "guessing is how a report ships with the previous project's table still in it")
        spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        try:
            build_from_template(Path(template_path), spec, Path(out_path), args.mode)
        except ValueError as e:
            print(f"ERROR: {e}", file=sys.stderr)
            sys.exit(1)
        print(f"-> {out_path}")
        return

    if args.revise:
        source_path, edits_path, out_path = args.revise
        data = json.loads(Path(edits_path).read_text(encoding="utf-8"))
        edits = data["edits"] if isinstance(data, dict) else data
        partial_path = _partial_output(Path(out_path), Path(source_path))
        try:
            applied, failures = revise(Path(source_path), edits, Path(out_path), partial_path)
        except ValueError as e:
            print(f"ERROR: {e}", file=sys.stderr)
            sys.exit(1)
        if failures:
            print("UNAPPLIED:", file=sys.stderr)
            for n, anchor, why in failures:
                print(f"  #{n}: {why}: {anchor!r}", file=sys.stderr)
            print(f"PARTIAL: applied {applied}/{len(edits)} edits; {len(failures)} unapplied. "
                  f"Incomplete result saved as {partial_path}, NOT {out_path}.", file=sys.stderr)
            sys.exit(2)
        print(f"applied {applied}/{len(edits)} edits -> {out_path}")
        return

    if args.validate:
        findings = validate(Path(args.validate), Path(args.source) if args.source else None)
        if findings:
            print(f"STRUCTURAL ISSUES ({len(findings)}):")
            for f in findings:
                print(f"  - {f}")
            sys.exit(1)
        print("structural validation: no issues found")
        return

    if args.render:
        docx_path, out_dir = args.render
        try:
            pdf_path = render(Path(docx_path), Path(out_dir))
        except (RuntimeError, subprocess.CalledProcessError) as e:
            print(f"ERROR: {e}", file=sys.stderr)
            sys.exit(2)
        print(pdf_path)


if __name__ == "__main__":
    main()
